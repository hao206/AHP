import os
import json
import uuid
import io
from typing import List, Dict, Any, Optional
from fastapi import FastAPI, HTTPException, UploadFile, File, Response
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from core.ahp_engine import (
    AHPMatrix, AHPHierarchy, run_hybrid_ahp_topsis, snap_to_saaty_scale, format_saaty_label,
    run_monte_carlo_ahp, aggregate_expert_matrices,
    compute_group_consensus, compute_fuzzy_ahp
)
from core.file_importer import (
    parse_uploaded_file, generate_sample_excel_template, generate_sample_csv_template
)

app = FastAPI(title="AHP Decision Studio Enterprise API", version="2.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

DATA_DIR = os.path.join(os.path.dirname(__file__), "data")
PROJECTS_FILE = os.path.join(DATA_DIR, "projects.json")
os.makedirs(DATA_DIR, exist_ok=True)

# ----------------- Models -----------------

class AutoTuneConsistencyRequest(BaseModel):
    elements: List[str]
    matrix: List[List[float]]
    target_cr: Optional[float] = 0.10

class GroupConsensusRequest(BaseModel):
    elements: List[str]
    expert_matrices: List[List[List[float]]]

class FuzzyAHPRequest(BaseModel):
    elements: List[str]
    fuzzy_matrix: List[List[List[float]]]

class MatrixEvaluationRequest(BaseModel):
    elements: List[str]
    matrix: List[List[float]]
    method: Optional[str] = "eigenvector"

class HierarchySynthesizeRequest(BaseModel):
    goal: str
    criteria: List[str]
    alternatives: List[str]
    criteria_matrix: List[List[float]]
    alt_matrices: Dict[str, List[List[float]]]
    method: Optional[str] = "eigenvector"

class DynamicSensitivityRequest(BaseModel):
    criteria: List[str]
    alternatives: List[str]
    criteria_weights: Dict[str, float]
    alt_matrices: Dict[str, List[List[float]]]

class GradientSensitivityRequest(BaseModel):
    goal: str
    criteria: List[str]
    alternatives: List[str]
    criteria_matrix: List[List[float]]
    alt_matrices: Dict[str, List[List[float]]]
    selected_criterion: str

class IncompleteCompleteRequest(BaseModel):
    elements: List[str]
    matrix: List[List[float]]

class HybridTopsisRequest(BaseModel):
    decision_matrix: List[List[float]]
    weights: List[float]
    criterion_types: List[str]
    alternatives: List[str]
    criteria: List[str]

class ProjectModel(BaseModel):
    id: Optional[str] = None
    title: str
    description: Optional[str] = ""
    goal: str
    criteria: List[str]
    alternatives: List[str]
    criteria_matrix: List[List[float]]
    alt_matrices: Dict[str, List[List[float]]]
    updated_at: Optional[str] = None

class MonteCarloRequest(BaseModel):
    goal: str
    criteria: List[str]
    alternatives: List[str]
    criteria_matrix: List[List[float]]
    alt_matrices: Dict[str, List[List[float]]]
    num_simulations: Optional[int] = 1000
    perturbation_pct: Optional[float] = 0.20

class GroupDecisionRequest(BaseModel):
    elements: List[str]
    expert_matrices: List[List[List[float]]]

# ----------------- Helper Functions -----------------

def get_default_templates() -> Dict[str, Any]:
    return {
        "template-vendor-selection": {
            "id": "template-vendor-selection",
            "title": "Enterprise ERP Solution Selection",
            "description": "Multi-criteria decision model for selecting the optimal Enterprise Resource Planning system.",
            "goal": "Select Optimal Enterprise ERP Platform",
            "criteria": ["Implementation Cost", "Business Functionality", "Vendor Reliability & SLA", "System Scalability"],
            "alternatives": ["SAP S/4HANA Enterprise", "Oracle Cloud ERP", "Odoo Enterprise Edition"],
            "criteria_matrix": [
                [1.0, 0.3333, 2.0, 0.5],
                [3.0, 1.0, 4.0, 2.0],
                [0.5, 0.25, 1.0, 0.3333],
                [2.0, 0.5, 3.0, 1.0]
            ],
            "alt_matrices": {
                "Implementation Cost": [
                    [1.0, 0.5, 0.2],
                    [2.0, 1.0, 0.25],
                    [5.0, 4.0, 1.0]
                ],
                "Business Functionality": [
                    [1.0, 2.0, 5.0],
                    [0.5, 1.0, 3.0],
                    [0.2, 0.3333, 1.0]
                ],
                "Vendor Reliability & SLA": [
                    [1.0, 1.5, 4.0],
                    [0.6667, 1.0, 3.0],
                    [0.25, 0.3333, 1.0]
                ],
                "System Scalability": [
                    [1.0, 1.0, 3.0],
                    [1.0, 1.0, 3.0],
                    [0.3333, 0.3333, 1.0]
                ]
            }
        },
        "template-car-purchase": {
            "id": "template-car-purchase",
            "title": "Executive Fleet Vehicle Procurement",
            "description": "Saaty's classic benchmark AHP model for corporate fleet vehicle selection.",
            "goal": "Procure Best Executive Fleet Sedan",
            "criteria": ["Acquisition Cost", "Safety Rating", "Fuel Economy & Eco", "Style & Comfort"],
            "alternatives": ["Toyota Camry Hybrid", "Honda Accord Executive", "Mazda 6 Signature"],
            "criteria_matrix": [
                [1.0, 0.3333, 3.0, 2.0],
                [3.0, 1.0, 5.0, 4.0],
                [0.3333, 0.2, 1.0, 0.5],
                [0.5, 0.25, 2.0, 1.0]
            ],
            "alt_matrices": {
                "Acquisition Cost": [
                    [1.0, 1.5, 0.6667],
                    [0.6667, 1.0, 0.5],
                    [1.5, 2.0, 1.0]
                ],
                "Safety Rating": [
                    [1.0, 1.0, 2.0],
                    [1.0, 1.0, 2.0],
                    [0.5, 0.5, 1.0]
                ],
                "Fuel Economy & Eco": [
                    [1.0, 2.0, 3.0],
                    [0.5, 1.0, 2.0],
                    [0.3333, 0.5, 1.0]
                ],
                "Style & Comfort": [
                    [1.0, 0.5, 0.3333],
                    [2.0, 1.0, 0.5],
                    [3.0, 2.0, 1.0]
                ]
            }
        },
        "template-cloud-provider": {
            "id": "template-cloud-provider",
            "title": "Global Cloud Infrastructure Assessment",
            "description": "Evaluation of hyper-scale cloud infrastructure for mission-critical workloads.",
            "goal": "Select Primary Public Cloud Infrastructure",
            "criteria": ["Global Latency & Uptime", "Service Pricing & TCO", "AI & Analytics Stack", "Enterprise Security"],
            "alternatives": ["Amazon Web Services (AWS)", "Microsoft Azure", "Google Cloud Platform (GCP)"],
            "criteria_matrix": [
                [1.0, 2.0, 0.5, 1.0],
                [0.5, 1.0, 0.3333, 0.5],
                [2.0, 3.0, 1.0, 2.0],
                [1.0, 2.0, 0.5, 1.0]
            ],
            "alt_matrices": {
                "Global Latency & Uptime": [
                    [1.0, 1.5, 1.2],
                    [0.6667, 1.0, 0.8],
                    [0.8333, 1.25, 1.0]
                ],
                "Service Pricing & TCO": [
                    [1.0, 0.8, 0.5],
                    [1.25, 1.0, 0.6667],
                    [2.0, 1.5, 1.0]
                ],
                "AI & Analytics Stack": [
                    [1.0, 0.5, 0.25],
                    [2.0, 1.0, 0.5],
                    [4.0, 2.0, 1.0]
                ],
                "Enterprise Security": [
                    [1.0, 1.0, 1.0],
                    [1.0, 1.0, 1.0],
                    [1.0, 1.0, 1.0]
                ]
            }
        }
    }

def load_projects_from_disk() -> Dict[str, Any]:
    if not os.path.exists(PROJECTS_FILE):
        default_projects = get_default_templates()
        save_projects_to_disk(default_projects)
        return default_projects
    try:
        with open(PROJECTS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return get_default_templates()

def save_projects_to_disk(data: Dict[str, Any]):
    with open(PROJECTS_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

# ----------------- Endpoints -----------------

@app.get("/api/health")
def health():
    return {"status": "ok", "service": "AHP Decision Studio Enterprise API", "version": "2.0.0"}

@app.post("/api/ahp/evaluate-matrix")
def evaluate_matrix(req: MatrixEvaluationRequest):
    try:
        matrix_obj = AHPMatrix(req.elements, req.matrix)
        res = matrix_obj.evaluate(req.method or "eigenvector")
        return res
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.post("/api/ahp/benchmark-methods")
def benchmark_methods(req: MatrixEvaluationRequest):
    try:
        matrix_obj = AHPMatrix(req.elements, req.matrix)
        bench = matrix_obj.benchmark_methods()
        return bench
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.post("/api/ahp/incomplete-complete")
def incomplete_complete(req: IncompleteCompleteRequest):
    try:
        matrix_obj = AHPMatrix(req.elements, req.matrix)
        completed_matrix = matrix_obj.complete_missing_comparisons()
        eval_res = matrix_obj.evaluate("eigenvector")
        return {
            "completed_matrix": completed_matrix.tolist(),
            "evaluation": eval_res
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.post("/api/ahp/synthesize")
def synthesize_hierarchy(req: HierarchySynthesizeRequest):
    try:
        hierarchy = AHPHierarchy(req.goal, req.criteria, req.alternatives)
        for i in range(len(req.criteria)):
            for j in range(len(req.criteria)):
                hierarchy.criteria_matrix.matrix[i, j] = req.criteria_matrix[i][j]
                
        for c in req.criteria:
            if c in req.alt_matrices:
                m = req.alt_matrices[c]
                for i in range(len(req.alternatives)):
                    for j in range(len(req.alternatives)):
                        hierarchy.alt_matrices[c].matrix[i, j] = m[i][j]
                        
        result = hierarchy.synthesize(req.method or "eigenvector")
        return result
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.post("/api/ahp/sensitivity/dynamic")
def dynamic_sensitivity(req: DynamicSensitivityRequest):
    try:
        hierarchy = AHPHierarchy("Sensitivity", req.criteria, req.alternatives)
        for c in req.criteria:
            if c in req.alt_matrices:
                m = req.alt_matrices[c]
                for i in range(len(req.alternatives)):
                    for j in range(len(req.alternatives)):
                        hierarchy.alt_matrices[c].matrix[i, j] = m[i][j]
                        
        # Local alt scores
        alt_local = []
        for a_idx, alt in enumerate(req.alternatives):
            score = sum(req.criteria_weights.get(c, 0.0) * hierarchy.alt_matrices[c].evaluate("eigenvector")["weights_list"][a_idx] for c in req.criteria)
            alt_local.append({"alternative": alt, "score": round(score, 4), "percentage": round(score * 100, 2)})
            
        alt_local.sort(key=lambda x: x["score"], reverse=True)
        for rank, item in enumerate(alt_local, 1):
            item["rank"] = rank
        return {"rankings": alt_local}
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.post("/api/ahp/gradient-sensitivity")
def gradient_sensitivity(req: GradientSensitivityRequest):
    try:
        hierarchy = AHPHierarchy(req.goal, req.criteria, req.alternatives)
        for i in range(len(req.criteria)):
            for j in range(len(req.criteria)):
                hierarchy.criteria_matrix.matrix[i, j] = req.criteria_matrix[i][j]
        for c in req.criteria:
            if c in req.alt_matrices:
                m = req.alt_matrices[c]
                for i in range(len(req.alternatives)):
                    for j in range(len(req.alternatives)):
                        hierarchy.alt_matrices[c].matrix[i, j] = m[i][j]

        res = hierarchy.compute_gradient_sensitivity(req.selected_criterion)
        return res
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.post("/api/ahp/hybrid-topsis")
def hybrid_topsis(req: HybridTopsisRequest):
    try:
        res = run_hybrid_ahp_topsis(
            req.decision_matrix,
            req.weights,
            req.criterion_types,
            req.alternatives,
            req.criteria
        )
        return res
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.post("/api/ahp/monte-carlo")
def monte_carlo(req: MonteCarloRequest):
    try:
        res = run_monte_carlo_ahp(
            criteria=req.criteria,
            alternatives=req.alternatives,
            criteria_matrix=req.criteria_matrix,
            alt_matrices=req.alt_matrices,
            num_simulations=req.num_simulations or 1000,
            perturbation_pct=req.perturbation_pct or 0.20
        )
        return res
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.post("/api/export/excel")
def export_excel(project: ProjectModel):
    """Generates an executive multi-tab Excel spreadsheet with formatting."""
    try:
        wb = openpyxl.Workbook()
        
        # Styles
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        accent_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        highlight_fill = PatternFill(start_color="06B6D4", end_color="06B6D4", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=16, bold=True, color="1E293B")
        bold_font = Font(name="Calibri", size=11, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        # Tab 1: Executive Summary & Rankings
        ws_rank = wb.active
        ws_rank.title = "Báo Cáo Tổng Hợp"
        ws_rank.views.sheetView[0].showGridLines = True

        ws_rank.append(["BÁO CÁO KẾT QUẢ QUYẾT ĐỊNH - AHP DECISION STUDIO"])
        ws_rank["A1"].font = title_font
        ws_rank.append([f"Dự án: {project.title}"])
        ws_rank.append([f"Mục tiêu: {project.goal}"])
        ws_rank.append([])

        # Synthesis
        hierarchy = AHPHierarchy(project.goal, project.criteria, project.alternatives)
        for i in range(len(project.criteria)):
            for j in range(len(project.criteria)):
                hierarchy.criteria_matrix.matrix[i, j] = project.criteria_matrix[i][j]
        for c in project.criteria:
            if c in project.alt_matrices:
                m = project.alt_matrices[c]
                for i in range(len(project.alternatives)):
                    for j in range(len(project.alternatives)):
                        hierarchy.alt_matrices[c].matrix[i, j] = m[i][j]
        synth = hierarchy.synthesize("eigenvector")

        ws_rank.append(["Thứ hạng", "Phương án", "Điểm số Trọng số Tổng hợp", "Tỷ lệ phần trăm", "Trạng thái Đề xuất"])
        for col_idx in range(1, 6):
            cell = ws_rank.cell(row=5, column=col_idx)
            cell.fill = accent_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for item in synth["rankings"]:
            row_num = ws_rank.max_row + 1
            ws_rank.append([
                item["rank"],
                item["alternative"],
                item["score"],
                f"{item['percentage']}%",
                "Lựa chọn Hàng đầu (Hạng 1)" if item["rank"] == 1 else ""
            ])
            for c_i in range(1, 6):
                cell = ws_rank.cell(row=row_num, column=c_i)
                cell.border = thin_border
                if item["rank"] == 1:
                    cell.font = bold_font

        ws_rank.append([])
        ws_rank.append(["Tỷ số Không Nhất quán Toàn cục (CR):", f"{synth['overall_consistency_ratio'] * 100:.2f}%"])
        ws_rank.append(["Đánh giá Nhất quán:", "HỢP LỆ (CR < 10%)" if synth["is_overall_consistent"] else "VƯỢT NGƯỠNG CHO PHÉP"])

        # Tab 2: Criteria Evaluation
        ws_crit = wb.create_sheet(title="Trọng Số Tiêu Chí")
        ws_crit.views.sheetView[0].showGridLines = True
        ws_crit.append(["ĐÁNH GIÁ TRỌNG SỐ TIÊU CHÍ & ĐỐI SÁNH PHƯƠNG PHÁP"])
        ws_crit["A1"].font = title_font
        ws_crit.append([])
        ws_crit.append(["Tiêu chí", "Véc-tơ riêng (EVM)", "Trung bình nhân (GMM)", "Trung bình Số học", "Tỷ lệ phần trăm"])
        for col_idx in range(1, 6):
            c_cell = ws_crit.cell(row=3, column=col_idx)
            c_cell.fill = header_fill
            c_cell.font = header_font

        crit_bench = hierarchy.criteria_matrix.benchmark_methods()
        for row_data in crit_bench["comparison"]:
            ws_crit.append([
                row_data["element"],
                row_data["evm_weight"],
                row_data["gmm_weight"],
                row_data["arith_weight"],
                f"{row_data['evm_pct']}%"
            ])

        # Auto-fit column widths
        for ws in [ws_rank, ws_crit]:
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"AHP_Decision_Studio_Report_{uuid.uuid4().hex[:6]}.xlsx"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/projects")
def list_projects():
    projects = load_projects_from_disk()
    return list(projects.values())

@app.post("/api/projects")
def save_project(proj: ProjectModel):
    projects = load_projects_from_disk()
    p_id = proj.id if proj.id else f"proj-{uuid.uuid4().hex[:8]}"
    proj_dict = proj.dict()
    proj_dict["id"] = p_id
    projects[p_id] = proj_dict
    save_projects_to_disk(projects)
    return proj_dict

@app.delete("/api/projects/{project_id}")
def delete_project(project_id: str):
    projects = load_projects_from_disk()
    if project_id in projects:
        del projects[project_id]
        save_projects_to_disk(projects)
        return {"status": "success", "message": f"Đã xóa dự án {project_id}"}
    raise HTTPException(status_code=404, detail="Không tìm thấy dự án.")

@app.get("/api/templates")
def list_templates():
    return list(get_default_templates().values())

# ----------------- File Import & Personalization Endpoints -----------------

@app.post("/api/import/file")
async def import_file(file: UploadFile = File(...)):
    """
    Intelligently parses uploaded Excel (.xlsx, .xls), CSV (.csv), or JSON (.json) files
    to personalize decision models with custom criteria, alternatives, and quantitative metrics.
    """
    try:
        content = await file.read()
        if not content:
            raise HTTPException(status_code=400, detail="Tệp tin tải lên rỗng.")
        
        parsed_project = parse_uploaded_file(content, file.filename)
        return parsed_project
    except ValueError as ve:
        raise HTTPException(status_code=400, detail=str(ve))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Lỗi phân tích cú pháp tệp tin: {str(e)}")

@app.get("/api/templates/download-sample-excel")
def download_sample_excel():
    """Generates and serves a formatted Excel template for user data entry."""
    try:
        excel_bytes = generate_sample_excel_template()
        return StreamingResponse(
            io.BytesIO(excel_bytes),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=AHP_Mau_Nhap_Lieu.xlsx"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/templates/download-sample-csv")
def download_sample_csv():
    """Generates and serves a clean CSV template for user data entry."""
    try:
        csv_str = generate_sample_csv_template()
        return Response(
            content=csv_str,
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=AHP_Mau_Nhap_Lieu.csv"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

# ----------------- Advanced Analytical Endpoints -----------------

@app.post("/api/ahp/group-decision")
def group_decision_consensus(req: GroupDecisionRequest):
    """
    Multi-Expert Group Decision Making (GDM):
    Aggregates multiple expert assessment matrices into a single consensus matrix
    using element-wise geometric mean (AIJ).
    """
    try:
        agg_matrix = aggregate_expert_matrices(req.expert_matrices)
        matrix_obj = AHPMatrix(req.elements, agg_matrix)
        eval_res = matrix_obj.evaluate("eigenvector")
        return {
            "aggregated_matrix": agg_matrix,
            "evaluation": eval_res
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.post("/api/ahp/auto-tune-consistency")
def auto_tune_consistency(req: AutoTuneConsistencyRequest):
    """
    Voracious-AHP Auto-Tuning Engine:
    Reduces inconsistency in pairwise matrices to achieve CR <= target_cr
    by adjusting the most inconsistent comparison pair.
    """
    try:
        matrix_obj = AHPMatrix(req.elements, req.matrix)
        result = matrix_obj.auto_reduce_inconsistency(target_cr=req.target_cr or 0.10)
        return result
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.post("/api/ahp/group-consensus")
def group_consensus_metric(req: GroupConsensusRequest):
    """
    Goepel's AHP Group Consensus Indicator S* (AHP-OS / Goepel 2013):
    Evaluates Shannon Beta entropy across multiple expert decision matrices.
    """
    try:
        return compute_group_consensus(req.expert_matrices, req.elements)
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.post("/api/ahp/fuzzy")
def evaluate_fuzzy_ahp(req: FuzzyAHPRequest):
    """
    Fuzzy AHP (Buckley's Geometric Mean Method) from pyDecision:
    Calculates Triangular Fuzzy Number (TFN) weights and crisp defuzzified priorities.
    """
    try:
        return compute_fuzzy_ahp(req.fuzzy_matrix, req.elements)
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

# Optional Single-Port Serving: Mount built frontend dist if available
frontend_dist_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "frontend", "dist"))
if os.path.exists(frontend_dist_path):
    from fastapi.staticfiles import StaticFiles
    app.mount("/", StaticFiles(directory=frontend_dist_path, html=True), name="static")


